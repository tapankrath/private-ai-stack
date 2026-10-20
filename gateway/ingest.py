"""Ingestion: files -> text -> chunks -> labels -> embeddings -> store.

- Watched folders, uploads, exported email (.eml/.mbox), notes (Markdown), PDFs, Word, HTML,
  code; images and scans through OCR (docling-serve) and audio through whisper.cpp when those
  services are enabled.
- Dedupe by SHA-256: unchanged files are skipped; changed files are re-indexed; removed files
  are hard-deleted (rows, vectors, keyword index).
- Secrets are redacted before storage; PII is tagged and raises the sensitivity label, never
  lowers it below the collection default.
"""

from __future__ import annotations

import email
import email.policy
import hashlib
import html.parser
import io
import mailbox
import mimetypes
import os
import re
import tempfile
import time
from pathlib import Path

import httpx
import yaml

from . import sens_level
from .redact import redact_secrets, sensitivity_for

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".log", ".org", ".tex"}
CODE_EXT = {".py", ".js", ".ts", ".tsx", ".jsx", ".go", ".rs", ".java", ".kt", ".c", ".h", ".cpp", ".hpp", ".cs", ".rb",
            ".php", ".sh", ".sql", ".swift", ".scala", ".lua", ".r", ".m", ".tf"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".ogg", ".flac", ".webm"}
OFFICE_EXT = {".pptx", ".xlsx"}
MAX_FILE_BYTES = 200 * 1024 * 1024
CHUNK_TOKENS = 800
OVERLAP = 0.15
CODE_DEF = re.compile(r"^(?:async\s+def|def|class|function|func|fn|pub\s+fn|impl|interface|struct|type|export\s+(?:default\s+)?"
                      r"(?:function|class|const)|public\s|private\s|protected\s|module|package)\b")


class _HTMLText(html.parser.HTMLParser):
    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip += 1
        elif tag in {"h1", "h2", "h3"}:
            self.out.append("\n" + "#" * int(tag[1]) + " ")
        elif tag in {"p", "br", "li", "div", "tr"}:
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def html_to_text(s: str) -> str:
    p = _HTMLText()
    p.feed(s)
    return re.sub(r"\n{3,}", "\n\n", "".join(p.out)).strip()


def approx_tokens(text: str) -> int:
    return int(len(text.split()) * 1.3) + 1


class Ingestor:
    def __init__(self, store, embedder, settings, docling_url: str | None = None, whisper_url: str | None = None,
                 transport: httpx.BaseTransport | None = None):
        self.store, self.embedder, self.settings = store, embedder, settings
        self.docling_url = docling_url or os.environ.get("GATEWAY_DOCLING_URL")
        self.docling_path = os.environ.get("GATEWAY_DOCLING_PATH", "/v1/convert/file")
        self.whisper_url = whisper_url or getattr(settings, "whisper_url", None)
        self.http = httpx.Client(timeout=600, transport=transport)

    # ------------------------------------------------------------------ parsing
    def parse(self, path: str, data: bytes) -> tuple[str, str, dict]:
        """Return (kind, text, meta). kind: markdown | code | skip."""
        ext = Path(path).suffix.lower()
        meta = {"mime": mimetypes.guess_type(path)[0] or "application/octet-stream", "title": Path(path).name}
        if ext in TEXT_EXT:
            return "markdown", data.decode("utf-8", errors="replace"), meta
        if ext in CODE_EXT:
            return "code", data.decode("utf-8", errors="replace"), meta
        if ext in {".html", ".htm"}:
            return "markdown", html_to_text(data.decode("utf-8", errors="replace")), meta
        if ext == ".eml":
            return "markdown", self._email(email.message_from_bytes(data, policy=email.policy.default), meta), meta
        if ext == ".mbox":
            return "markdown", self._mbox(data, meta), meta
        if ext == ".pdf":
            text = self._pdf(data)
            if len(text.strip()) < 50 * max(1, text.count("## Page")) and self.docling_url:
                text = self._docling(path, data) or text          # scanned PDF -> OCR
            return "markdown", text, meta
        if ext == ".docx":
            return "markdown", self._docx(data, meta), meta
        if ext in IMAGE_EXT or ext in OFFICE_EXT:
            if not self.docling_url:
                return "skip", "", {**meta, "reason": "needs the OCR/document service (docling) which is not enabled"}
            return "markdown", self._docling(path, data) or "", meta
        if ext in AUDIO_EXT:
            if not self.whisper_url:
                return "skip", "", {**meta, "reason": "needs speech-to-text (whisper.cpp) which is not enabled"}
            return "markdown", self._whisper(path, data), meta
        return "skip", "", {**meta, "reason": f"unsupported file type {ext or '(none)'}"}

    def _email(self, msg, meta) -> str:
        subj = str(msg.get("subject", "(no subject)"))
        meta["title"] = subj
        meta["author"] = str(msg.get("from", ""))
        try:
            meta["created_at"] = email.utils.parsedate_to_datetime(msg.get("date")).timestamp()
        except (TypeError, ValueError):
            pass
        body = msg.get_body(preferencelist=("plain", "html"))
        text = ""
        if body is not None:
            text = body.get_content()
            if body.get_content_type() == "text/html":
                text = html_to_text(text)
        atts = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
        head = f"## {subj}\nFrom: {msg.get('from', '')}\nTo: {msg.get('to', '')}\nDate: {msg.get('date', '')}\n"
        if atts:
            head += f"Attachments: {', '.join(atts)}\n"
        return head + "\n" + text

    def _mbox(self, data: bytes, meta) -> str:
        with tempfile.NamedTemporaryFile(suffix=".mbox", delete=False) as f:
            f.write(data)
            name = f.name
        try:
            box = mailbox.mbox(name, factory=lambda fh: email.message_from_binary_file(fh, policy=email.policy.default))
            parts = [self._email(m, {}) for m in box]
        finally:
            os.unlink(name)
        meta["title"] = f"{meta['title']} ({len(parts)} emails)"
        return "\n\n".join(parts)

    def _pdf(self, data: bytes) -> str:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(f"## Page {i}\n{(p.extract_text() or '').strip()}" for i, p in enumerate(reader.pages, 1))

    def _docx(self, data: bytes, meta) -> str:
        import docx

        d = docx.Document(io.BytesIO(data))
        meta["author"] = d.core_properties.author or None
        out = []
        for p in d.paragraphs:
            style = (p.style.name or "").lower() if p.style is not None else ""
            m = re.match(r"heading (\d)", style)
            if m and p.text.strip():
                out.append("#" * int(m.group(1)) + " " + p.text.strip())
            elif p.text.strip():
                out.append(p.text.strip())
        for t in d.tables:
            for row in t.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n\n".join(out)

    def _docling(self, path: str, data: bytes) -> str | None:
        try:
            r = self.http.post(self.docling_url.rstrip("/") + self.docling_path,
                               files={"files": (Path(path).name, data)}, data={"to_formats": "md", "do_ocr": "true"})
            r.raise_for_status()
            j = r.json()
            doc = j.get("document") or (j.get("documents") or [{}])[0]
            return doc.get("md_content") or doc.get("markdown")
        except (httpx.HTTPError, ValueError, KeyError, IndexError):
            return None

    def _whisper(self, path: str, data: bytes) -> str:
        r = self.http.post(self.whisper_url.rstrip("/") + "/inference", files={"file": (Path(path).name, data)},
                           data={"response_format": "json"})
        r.raise_for_status()
        return "## Transcript\n" + r.json().get("text", "")

    # ------------------------------------------------------------------ chunking
    def chunk(self, kind: str, text: str, path: str) -> list[dict]:
        return self._chunk_code(text, path) if kind == "code" else self._chunk_markdown(text)

    def _pack(self, units: list[tuple[str, str]]) -> list[dict]:
        """units: (heading_path, paragraph). Pack into ~CHUNK_TOKENS with ~15% overlap within a heading."""
        chunks, cur, cur_head, cur_tok = [], [], None, 0
        for head, para in units:
            t = approx_tokens(para)
            if cur and (head != cur_head or cur_tok + t > CHUNK_TOKENS):
                chunks.append((cur_head, "\n\n".join(cur)))
                if head == cur_head:                       # overlap only inside the same section
                    keep, kt = [], 0
                    for p in reversed(cur):
                        kt += approx_tokens(p)
                        if kt > CHUNK_TOKENS * OVERLAP:
                            break
                        keep.insert(0, p)
                    cur, cur_tok = keep, sum(approx_tokens(p) for p in keep)
                else:
                    cur, cur_tok = [], 0
            if t > CHUNK_TOKENS:                           # very long paragraph: split by words
                words = para.split()
                step = int(CHUNK_TOKENS / 1.3)
                for i in range(0, len(words), int(step * (1 - OVERLAP))):
                    chunks.append((head, " ".join(words[i:i + step])))
                cur, cur_tok, cur_head = [], 0, head
                continue
            cur.append(para)
            cur_tok += t
            cur_head = head
        if cur:
            chunks.append((cur_head, "\n\n".join(cur)))
        return [{"ord": i, "heading_path": h or "", "text": t, "tokens": approx_tokens(t)} for i, (h, t) in enumerate(chunks) if t.strip()]

    def _chunk_markdown(self, text: str) -> list[dict]:
        units, stack, para = [], [], []

        def flush():
            if para:
                units.append((" > ".join(stack), "\n".join(para).strip()))
                para.clear()

        for line in text.splitlines():
            m = re.match(r"^(#{1,6})\s+(.*)$", line)
            if m:
                flush()
                level = len(m.group(1))
                stack[:] = stack[: level - 1] + [m.group(2).strip()]
            elif not line.strip():
                flush()
            else:
                para.append(line)
        flush()
        return self._pack([u for u in units if u[1]])

    def _chunk_code(self, text: str, path: str) -> list[dict]:
        units, cur, name = [], [], Path(path).name
        for line in text.splitlines():
            if CODE_DEF.match(line) and cur:
                units.append((name, "\n".join(cur)))
                cur = []
            if CODE_DEF.match(line):
                sym = re.sub(r"[\s(:{].*$", "", re.sub(CODE_DEF, "", line).strip()) or line.strip()[:40]
                name = f"{Path(path).name} > {sym}"
            cur.append(line)
        if cur:
            units.append((name, "\n".join(cur)))
        return self._pack([(h, t) for h, t in units if t.strip()])

    # ------------------------------------------------------------------ ingest
    def ingest_bytes(self, collection: dict, path: str, data: bytes, mtime: float | None = None) -> dict:
        if len(data) > MAX_FILE_BYTES:
            return {"status": "skipped", "reason": "file too large"}
        sha = hashlib.sha256(data).hexdigest()
        existing = self.store.get_document(collection["id"], path)
        if existing and existing["sha256"] == sha:
            return {"status": "unchanged", "chunks": 0}
        kind, text, meta = self.parse(path, data)
        if kind == "skip":
            return {"status": "skipped", "reason": meta.get("reason")}
        meta.setdefault("modified_at", mtime or time.time())
        floor = int(collection["default_sensitivity"])
        chunks = self.chunk(kind, text, path)
        doc_sens, doc_pii = floor, set()
        for c in chunks:
            red, _ = redact_secrets(c["text"])
            sens, tags = sensitivity_for(c["text"], floor)       # label from the original, store the redacted
            c["text"], c["sensitivity"], c["pii"] = red, sens, tags
            doc_sens, doc_pii = max(doc_sens, sens), doc_pii | set(tags)
        try:
            vecs = self.embedder.embed([(c["heading_path"] + "\n" + c["text"]).strip() for c in chunks]) if chunks else []
        except Exception:
            vecs = [None] * len(chunks)
        self.store.upsert_document(collection["id"], path, sha, meta, doc_sens, sorted(doc_pii), chunks, vecs)
        return {"status": "updated" if existing else "added", "chunks": len(chunks), "sensitivity": doc_sens,
                "pii": sorted(doc_pii)}

    def scan_folder(self, collection: dict, folder: str) -> dict:
        root = Path(folder)
        seen, counts = set(), {"added": 0, "updated": 0, "unchanged": 0, "skipped": 0, "removed": 0, "errors": 0}
        for p in sorted(root.rglob("*")):
            if not p.is_file() or any(part.startswith(".") for part in p.relative_to(root).parts):
                continue
            rel = str(p.relative_to(root))
            seen.add(rel)
            try:
                res = self.ingest_bytes(collection, rel, p.read_bytes(), p.stat().st_mtime)
                counts[res["status"]] += 1
            except Exception:  # one bad file never stops the scan
                counts["errors"] += 1
        for d in self.store.documents_in(collection["id"]):
            if d["path"] not in seen:
                self.store.delete_document(collection["id"], d["path"])
                counts["removed"] += 1
        return counts


def load_collections(path: str) -> list[dict]:
    """collections.yaml: [{name, path, sensitivity, retention_days, owner, purpose, grants: [{principal, permission}]}]"""
    cfg = yaml.safe_load(Path(path).read_text()) or {}
    return cfg.get("collections", [])


def sync_collections(store, cfg: list[dict]) -> list[tuple[dict, str]]:
    out = []
    for c in cfg:
        owner = c.get("owner", "owner")
        if not store.get_user(owner):
            store.upsert_user(owner, owner, "member", [])
        cid = store.create_collection(c["name"], owner, sens_level(c.get("sensitivity", "personal")),
                                      c.get("retention_days"), c.get("purpose", ""))
        for g in c.get("grants", []):
            store.grant(cid, g["principal"], g.get("permission", "read"))
        out.append((store.get_collection(cid), c["path"]))
    return out
